"""Downloadable inspection report (PDF, DOCX, TXT) built only from stored records."""

import io
from datetime import UTC, datetime

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF
from PIL import Image

FORMATS = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
}
DECISIONS = {"normal": "Pass (normal)", "anomaly": "Fail (anomaly)", "review": "Needs review"}
LABELS = {"normal": "Pass", "anomaly": "Fail", "unsure": "Unsure"}


def collect(service, inspection_id):
    """Gather the inspection, its profile, engineer feedback and latest AI description."""
    inspection = service.store.get("inspection", inspection_id)
    if inspection.get("status") != "ready":
        raise ValueError("Inspection has not completed")
    profile = service.store.get("profile", inspection["profile_id"])
    versions = sorted(
        (v for v in service.store.all("version") if v["profile_id"] == profile["id"]),
        key=lambda v: v["created_at"],
    )
    number = next(
        (i + 1 for i, v in enumerate(versions) if v["id"] == inspection["version_id"]), None
    )
    feedback = next(
        (f for f in service.store.all("feedback") if f["inspection_id"] == inspection_id), None
    )
    try:
        report = service.latest_report(inspection_id)
    except KeyError:
        report = None
    asset = lambda url: url.rsplit("/", 1)[-1] if url else None  # noqa: E731
    images = [
        ("Inspected part", inspection["image"]["asset"]),
        ("Heatmap overlay", asset(inspection.get("overlay_url"))),
        ("Closest good reference", asset(inspection.get("reference_url"))),
    ]
    return {
        "inspection": inspection,
        "profile": profile,
        "version_number": number,
        "feedback": feedback,
        "report": report if report and report.get("status") == "ready" else None,
        "images": [(caption, name) for caption, name in images if name],
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
    }


def _lines(data):
    """The report as (kind, text) rows shared by all formats; kinds: h1, h2, kv, p, li."""
    i, profile, runtime = data["inspection"], data["profile"], data["inspection"].get("runtime", {})
    device = runtime.get("device_name") or runtime.get("device") or "unknown"
    rows = [
        ("h1", "Worksight inspection report"),
        ("kv", ("Generated", data["generated_at"])),
        ("kv", ("Inspection", i["id"])),
        ("kv", ("Inspected at", i["created_at"])),
        ("kv", ("Profile", f"{profile['name']} (version {data['version_number'] or '?'})")),
        ("kv", ("File", i.get("filename") or "uploaded image")),
        ("images", None),
        ("h2", "Model result"),
        ("kv", ("Decision", DECISIONS.get(i["decision"], i["decision"]))),
        ("kv", ("Anomaly score", f"{i['anomaly_score']:.4f}")),
        ("kv", ("Threshold", f"{i['threshold']:.4f}")),
        ("kv", ("Inference time", f"{i['inference_ms']:.0f} ms (preprocessing, encoding, retrieval)")),
        ("kv", ("Device", f"{device}{' (AMD GPU)' if runtime.get('amd_gpu') else ''}")),
        ("p", "The anomaly score is a distance to the good references, not a probability."),
        ("h2", "Engineer review"),
    ]
    feedback = data["feedback"]
    if feedback:
        rows += [
            ("kv", ("Decision", LABELS.get(feedback["label"], feedback["label"]))),
            ("kv", ("Recorded", feedback["created_at"])),
            ("kv", ("Note", feedback.get("note") or "—")),
        ]
    else:
        rows.append(("p", "Not reviewed yet."))
    rows.append(("h2", "AI description (assistive)"))
    report = data["report"]
    if report:
        rows += [
            ("kv", ("Model", f"{report.get('model')} on AMD Token Factory")),
            ("p", report.get("summary") or ""),
            *(("li", o) for o in report.get("observations") or []),
        ]
        if report.get("suggested_check"):
            rows.append(("kv", ("Check first", report["suggested_check"])))
        for requirement in report.get("relevant_requirements") or []:
            rows.append(("li", f"Requirement: {requirement}"))
        rows.append(("p", "Generated text can be wrong; it does not change the model or engineer decision."))
    else:
        rows.append(("p", "No AI description was generated for this inspection."))
    return rows


def to_txt(data) -> bytes:
    out = []
    for kind, value in _lines(data):
        if kind == "h1":
            out += [value, "=" * len(value)]
        elif kind == "h2":
            out += ["", value, "-" * len(value)]
        elif kind == "kv":
            out.append(f"{value[0]}: {value[1]}")
        elif kind == "li":
            out.append(f"- {value}")
        elif kind == "images":
            out.append("Images: " + ", ".join(caption for caption, _ in data["images"]))
        else:
            out.append(value)
    return ("\n".join(out) + "\n").encode()


def _latin1(text: str) -> str:
    # Core PDF fonts are Latin-1; keep the meaning of common Unicode punctuation.
    for a, b in {"–": "-", "—": "-", "·": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "…": "..."}.items():
        text = text.replace(a, b)
    return text.encode("latin-1", "replace").decode("latin-1")


def _jpeg(image: Image.Image, side=900) -> io.BytesIO:
    # JPEG keeps a 3-image report around a few hundred KB instead of several MB.
    image = image.convert("RGB")
    image.thumbnail((side, side))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=85)
    buffer.seek(0)
    return buffer


def _pdf_images(pdf, data, load_image, width):
    # The three images side by side, so the whole report fits on one page.
    images = [(caption, load_image(name)) for caption, name in data["images"]]
    if not images:
        return
    gap = 4
    w = (width - gap * (len(images) - 1)) / len(images)
    h = max(w * image.height / image.width for _, image in images)
    top = pdf.get_y() + 2
    for n, (caption, image) in enumerate(images):
        x = pdf.l_margin + n * (w + gap)
        pdf.set_xy(x, top)
        pdf.set_font("Helvetica", "", 8)
        pdf.cell(w, 4, _latin1(caption))
        pdf.image(_jpeg(image), x=x, y=top + 5, w=w, h=w * image.height / image.width)
    pdf.set_xy(pdf.l_margin, top + 5 + h + 2)


def to_pdf(data, load_image) -> bytes:
    pdf = FPDF()
    pdf.set_margins(14, 12, 14)
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    for kind, value in _lines(data):
        if kind == "h1":
            pdf.set_font("Helvetica", "B", 16)
            pdf.multi_cell(width, 8, _latin1(value), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif kind == "h2":
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(width, 6, _latin1(value), new_x="LMARGIN", new_y="NEXT")
            pdf.set_draw_color(220, 220, 220)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + width, pdf.get_y())
            pdf.ln(1)
        elif kind == "kv":
            pdf.set_font("Helvetica", "B", 9)
            pdf.cell(36, 5, _latin1(value[0]))
            pdf.set_font("Helvetica", "", 9)
            pdf.multi_cell(width - 36, 5, _latin1(str(value[1])), new_x="LMARGIN", new_y="NEXT")
        elif kind == "li":
            pdf.set_font("Helvetica", "", 9)
            pdf.multi_cell(width, 5, _latin1(f"- {value}"), new_x="LMARGIN", new_y="NEXT")
        elif kind == "images":
            _pdf_images(pdf, data, load_image, width)
        else:
            pdf.set_font("Helvetica", "", 9)
            pdf.multi_cell(width, 5, _latin1(value), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def to_docx(data, load_image) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for kind, value in _lines(data):
        if kind == "h1":
            doc.add_heading(value, level=0)
        elif kind == "h2":
            doc.add_heading(value, level=1)
        elif kind == "kv":
            p = doc.add_paragraph()
            p.add_run(f"{value[0]}: ").bold = True
            p.add_run(str(value[1]))
        elif kind == "li":
            doc.add_paragraph(value, style="List Bullet")
        elif kind == "images" and data["images"]:
            table = doc.add_table(rows=2, cols=len(data["images"]))
            for n, (caption, name) in enumerate(data["images"]):
                table.cell(0, n).paragraphs[0].add_run(caption).bold = True
                run = table.cell(1, n).paragraphs[0].add_run()
                run.add_picture(_jpeg(load_image(name)), width=Inches(2.0))
        elif kind != "images":
            doc.add_paragraph(value)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render(service, inspection_id, fmt):
    if fmt not in FORMATS:
        raise ValueError("Format must be pdf, docx or txt")
    data = collect(service, inspection_id)
    if fmt == "txt":
        body = to_txt(data)
    elif fmt == "pdf":
        body = to_pdf(data, service.store.image)
    else:
        body = to_docx(data, service.store.image)
    return body, FORMATS[fmt], f"worksight-inspection-{inspection_id[:8]}.{fmt}"
